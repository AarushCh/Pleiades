import argparse
import hashlib
import json
import os
import random
import re
import subprocess
import sys
from datetime import datetime, timezone
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
HERE = Path(__file__).resolve().parent
BUILD = HERE / "build"
CORPUS = BUILD / "corpus"
SEED = HERE / "seed_qa.jsonl"

SECRET = re.compile(
    r"gsk_[A-Za-z0-9]{20,}|sk-[A-Za-z0-9_-]{20,}|AKIA[0-9A-Z]{16}"
    r"|-----BEGIN [A-Z ]*PRIVATE KEY-----|postgres(?:ql)?(?:\+\w+)?://[^\s:/]+:[^\s@]+@"
)
NUMBER = re.compile(r"\d[\d,]*(?:\.\d+)?")
TICKS = re.compile(r"`([^`]+)`")
REFUSAL = re.compile(r"does not cover|doesn't cover|not covered|no information|not in the (project )?knowledge base", re.I)
REFUSAL_TEXT = ("The project knowledge base does not cover that. Ask me about building Pleiades instead: its code, "
                "data, design decisions, security, deployment, roadmap or training.")

COPIES = {"project-summary.md": "Pleaides_Summary.md", "readme.md": "README.md"}
GROUPS = {
    "code-backend.md": ("Backend source code", ["src/*.py", "api/*.py"]),
    "code-frontend.md": ("Frontend source code", ["frontend/app/*.jsx", "frontend/components/*.jsx", "frontend/lib/*.js",
                                                  "frontend/next.config.mjs", "frontend/package.json"]),
    "code-tests-eval.md": ("Tests and evaluation", ["tests/*.py", "eval/run_eval.py", "eval/dataset.json"]),
    "code-training.md": ("Training kit", ["training/pipeline.py", "training/finetune.py", "training/assistant_prompt.md",
                                          "training/Modelfile"]),
    "ops-config.md": ("Deployment and configuration", ["Dockerfile", "render.yaml", ".env.example", "requirements.txt",
                                                      "run.ps1", "setup.ps1", "setup.sh", ".gitignore", "pytest.ini"]),
}
LANG = {".py": "python", ".jsx": "jsx", ".js": "javascript", ".mjs": "javascript", ".json": "json",
        ".yaml": "yaml", ".ps1": "powershell", ".sh": "bash", ".md": "markdown"}

GEN_SYSTEM = (
    "You write training data for the Pleiades project assistant. From the passage, write {n} question and answer "
    "pairs that a team member building Pleiades might ask. Every answer must be fully supported by the passage, "
    "1 to 4 sentences, and must copy numbers, names, commands, environment variables and file paths exactly. "
    "Mix fact, how, why, which-file and what-command questions. Return only a JSON array of objects with the keys "
    "question and answer."
)

OFF_TOPIC = [
    "Who won the cricket world cup in 2011?", "What is the weather in Visakhapatnam today?",
    "Write me a poem about the monsoon.", "What is the capital of Australia?",
    "How do I cook biryani?", "What is the share price of Reliance today?",
    "Who is the prime minister of Japan?", "Recommend a good movie for tonight.",
    "How many moons does Jupiter have?", "Translate 'good morning' into French.",
    "What is the best smartphone to buy this year?", "Explain the theory of relativity.",
    "How do I file my income tax return?", "What time does the railway ticket counter open?",
    "Who wrote the Ramayana?", "What are the symptoms of dengue?",
    "Suggest a workout plan for beginners.", "How tall is Mount Everest?",
    "What is the exchange rate from dollars to rupees?", "How do I reset my Instagram password?",
]


def numbers(text):
    return {n.replace(",", "").rstrip(".") for n in NUMBER.findall(text)}


def norm(text):
    return " ".join(re.findall(r"[a-z0-9]+", text.lower()))


def verified_pairs(raw, passage):
    start, end = raw.find("["), raw.rfind("]")
    if start == -1 or end <= start:
        return []
    try:
        items = json.loads(raw[start:end + 1])
    except json.JSONDecodeError:
        return []
    allowed, kept = numbers(passage), []
    for item in items if isinstance(items, list) else []:
        if not isinstance(item, dict):
            continue
        q, a = str(item.get("question", "")).strip(), str(item.get("answer", "")).strip()
        if not (10 <= len(q) <= 300 and 20 <= len(a) <= 900):
            continue
        if not numbers(f"{q} {a}") <= allowed:
            continue
        if any(t not in passage for t in TICKS.findall(a)):
            continue
        kept.append({"question": q, "answer": a})
    return kept


def split_of(key):
    return "val" if int(hashlib.sha256(key.encode()).hexdigest()[:8], 16) % 10 == 0 else "train"


def seed_items():
    return [json.loads(line) for line in SEED.read_text(encoding="utf-8").splitlines() if line.strip()]


def commits():
    raw = subprocess.run(["git", "log", "--date=short", "--format=%h|%ad|%an|%s%n%b%x1e"], cwd=ROOT,
                         capture_output=True, text=True, encoding="utf-8").stdout
    out = ["# Commit history", ""]
    for entry in (e.strip() for e in raw.split("\x1e")):
        if not entry:
            continue
        head, _, body = entry.partition("\n")
        h, date, author, subject = head.split("|", 3)
        out += [f"## {subject}", "", f"Commit {h} on {date} by {author}.", "", body.strip(), ""]
    return "\n".join(out)


def deck():
    path = ROOT / "CapstoneStuff" / "Pleiades_Review1.pptx"
    try:
        from pptx import Presentation
    except ImportError:
        return None
    if not path.exists():
        return None
    out = ["# Review 1 presentation", ""]
    for i, slide in enumerate(Presentation(path).slides, 1):
        lines = []
        for sh in slide.shapes:
            if sh.name.startswith("Footer"):
                continue
            if sh.has_text_frame and sh.text_frame.text.strip():
                lines.append(" ".join(sh.text_frame.text.split()))
            elif getattr(sh, "has_table", False) and sh.has_table:
                lines += [" | ".join(c.text for c in row.cells) for row in sh.table.rows]
        title = lines[0] if lines else f"Slide {i}"
        out += [f"## Slide {i}: {title}", ""] + [f"- {ln}" for ln in lines[1:]]
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            out += ["", f"Speaker notes: {slide.notes_slide.notes_text_frame.text.strip()}"]
        out.append("")
    return "\n".join(out)


def cmd_build(_):
    CORPUS.mkdir(parents=True, exist_ok=True)
    for old in CORPUS.glob("*.md"):
        old.unlink()
    docs = {name: (ROOT / src).read_text(encoding="utf-8") for name, src in COPIES.items() if (ROOT / src).exists()}
    docs.update({p.name: p.read_text(encoding="utf-8") for p in sorted((ROOT / "data").glob("*.md"))})
    for name, (title, patterns) in GROUPS.items():
        parts = [f"# {title}", ""]
        for pattern in patterns:
            for p in sorted(ROOT.glob(pattern)):
                parts += [f"## {p.relative_to(ROOT).as_posix()}", "", f"```{LANG.get(p.suffix, 'text')}",
                          p.read_text(encoding="utf-8").rstrip(), "```", ""]
        docs[name] = "\n".join(parts)
    docs["commit-history.md"] = commits()
    slides = deck()
    if slides:
        docs["review1-deck.md"] = slides
    manifest, redacted = [], 0
    for name, body in docs.items():
        body, n = SECRET.subn("[REDACTED]", body)
        redacted += n
        (CORPUS / name).write_text(body, encoding="utf-8")
        manifest.append({"file": name, "chars": len(body), "sha256": hashlib.sha256(body.encode()).hexdigest()})
    (BUILD / "manifest.json").write_text(json.dumps({"built": datetime.now(timezone.utc).isoformat(),
                                                     "files": manifest}, indent=2), encoding="utf-8")
    print(f"Corpus: {len(docs)} documents, {sum(m['chars'] for m in manifest):,} characters, {redacted} secrets redacted")
    from src.ingest import build_index
    build_index(rebuild=True)


def assistant():
    from src import config
    from src.rag import SupportAssistant
    if not config.CHROMA_DIR.exists():
        raise SystemExit("No project index. Run: python training/pipeline.py build")
    return SupportAssistant()


def cmd_dataset(args):
    from langchain_core.output_parsers import StrOutputParser
    from langchain_core.prompts import ChatPromptTemplate

    from src import config
    from src.llm import build_llm, resolve_backend
    from src.rag import SYSTEM_PROMPT, format_context

    config.EXPAND_THRESHOLD = 0.0
    bot = assistant()
    random.seed(args.seed)
    pairs_path = BUILD / "qa_pairs.jsonl"

    if not args.assemble_only:
        backend = resolve_backend()
        if backend == "stub":
            raise SystemExit("Generating pairs needs a language model: set LLAMA_API_* in .env or run Ollama")
        llm, name = build_llm(backend)
        chain = ChatPromptTemplate.from_messages([("system", GEN_SYSTEM),
                                                  ("human", "Source: {label} > {section}\n\n{passage}")]) | llm | StrOutputParser()
        done = {json.loads(line)["chunk"] for line in pairs_path.read_text(encoding="utf-8").splitlines()} \
            if pairs_path.exists() else set()
        todo = [c for c in bot.chunks if args.only in c.metadata["filename"] and c.metadata["hash"] not in done
                and len(c.page_content) >= 200]
        if args.limit:
            todo = todo[:args.limit]
        print(f"Generating with {name}: {len(todo)} chunks, {args.per_chunk} pairs each")
        with pairs_path.open("a", encoding="utf-8") as f:
            for i, c in enumerate(todo, 1):
                section = c.metadata.get("section") or c.metadata.get("doc_title") or ""
                try:
                    raw = chain.invoke({"n": args.per_chunk, "label": c.metadata["source_label"],
                                        "section": section, "passage": c.page_content})
                except Exception as exc:
                    print(f"  [{i}/{len(todo)}] model error: {type(exc).__name__}: {str(exc)[:120]}")
                    continue
                kept = verified_pairs(raw, c.page_content)
                f.write(json.dumps({"chunk": c.metadata["hash"], "pairs": kept}, ensure_ascii=False) + "\n")
                print(f"  [{i}/{len(todo)}] {c.metadata['source_label']} > {section[:48]}: kept {len(kept)}")

    held_out = {norm(it["question"]) for it in seed_items()}
    records = [json.loads(line) for line in pairs_path.read_text(encoding="utf-8").splitlines() if line.strip()] \
        if pairs_path.exists() else []
    seen, splits, counts = set(), {"train": [], "val": []}, {"grounded": 0, "refusal": 0, "held_out_skipped": 0}

    def sample(question, answer, docs, split, kind):
        splits[split].append({"messages": [
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": f"<context>\n{format_context(docs)}\n</context>\n<question>{question}</question>"},
            {"role": "assistant", "content": answer}], "kind": kind})
        counts[kind] += 1

    for rec in records:
        chunk = bot.by_hash.get(rec["chunk"])
        if chunk is None:
            continue
        label = chunk.metadata["source_label"]
        section = chunk.metadata.get("section") or chunk.metadata.get("doc_title") or ""
        for pair in rec["pairs"]:
            key = norm(pair["question"])
            if key in held_out:
                counts["held_out_skipped"] += 1
                continue
            if key in seen:
                continue
            seen.add(key)
            docs = [d for d in bot.retrieve(pair["question"], []).docs if d.metadata["hash"] != chunk.metadata["hash"]]
            docs = docs[:bot.top_k - 1]
            docs.insert(random.randrange(len(docs) + 1), chunk)
            src = f"{label} > {section}" if section else label
            sample(pair["question"], f"{pair['answer']}\n\nSources: {src}", docs, split_of(f"{label}|{section}"), "grounded")

    for _ in range(int(counts["grounded"] * args.refusals)):
        question = random.choice(OFF_TOPIC)
        sample(question, REFUSAL_TEXT, bot.retrieve(question, []).docs, split_of(question), "refusal")

    for split, rows in splits.items():
        random.shuffle(rows)
        with (BUILD / f"sft_{split}.jsonl").open("w", encoding="utf-8") as f:
            for row in rows:
                f.write(json.dumps(row, ensure_ascii=False) + "\n")
    print(f"SFT set: {len(splits['train'])} train, {len(splits['val'])} val, {counts}")


def cmd_eval(args):
    from src import config
    if not args.answers:
        config.EXPAND_THRESHOLD = 0.0
    bot = assistant()
    if args.answers and not bot.has_llm:
        raise SystemExit("--answers needs a language model: set LLAMA_API_* in .env or run Ollama")
    rows = []
    for it in seed_items():
        r = bot.retrieve(it["question"], [])
        labels = {d.metadata["source_label"] for d in r.docs}
        row = {"id": it["id"], "hit": None if it.get("refuse") else bool(labels & set(it["source"]))}
        if args.answers:
            payload = bot._payload(it["question"], r, [])
            answer = bot.answer_chain.invoke(payload).strip()
            refused = bool(REFUSAL.search(answer))
            row.update(answer=answer, refused=refused,
                       grounded=numbers(answer) <= numbers(payload["context"] + it["question"]),
                       facts=all(k.lower() in answer.lower() for k in it.get("must", [])))
        rows.append(row)
        mark = "-" if row["hit"] is None else ("hit " if row["hit"] else "MISS")
        print(f"  {mark:4} {it['id']}" + (f"  facts={row['facts']} grounded={row['grounded']} refused={row['refused']}"
                                          if args.answers else ""))
    scored = [r for r in rows if r["hit"] is not None]
    summary = {"model": bot.backend_name, "items": len(rows),
               "retrieval_hit_rate": round(sum(r["hit"] for r in scored) / len(scored), 3)}
    if args.answers:
        refuse_ids = {it["id"] for it in seed_items() if it.get("refuse")}
        answered = [r for r in rows if r["id"] not in refuse_ids]
        summary.update(
            fact_pass_rate=round(sum(r["facts"] for r in answered) / len(answered), 3),
            grounded_rate=round(sum(r["grounded"] for r in rows) / len(rows), 3),
            false_refusals=sum(r["refused"] for r in answered),
            correct_refusals=f"{sum(r['refused'] for r in rows if r['id'] in refuse_ids)}/{len(refuse_ids)}")
    out = BUILD / f"eval_{datetime.now().strftime('%Y%m%d_%H%M%S')}.json"
    out.write_text(json.dumps({"summary": summary, "rows": rows}, indent=2, ensure_ascii=False), encoding="utf-8")
    print(json.dumps(summary, indent=2))


def cmd_chat(args):
    from src import cli
    sys.argv = [sys.argv[0], *args.question]
    cli.main()


def cmd_serve(args):
    import uvicorn
    uvicorn.run("api.main:app", host="127.0.0.1", port=args.port)


def main():
    os.environ.setdefault("KB_DIR", str(CORPUS))
    os.environ.setdefault("CHROMA_DIR", str(BUILD / "chroma"))
    os.environ.setdefault("COLLECTION_NAME", "pleiades_project")
    os.environ.setdefault("KB_FROM_DB", "0")
    os.environ.setdefault("SYSTEM_PROMPT_FILE", str(HERE / "assistant_prompt.md"))
    sys.path.insert(0, str(ROOT))
    if hasattr(sys.stdout, "reconfigure"):
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")

    parser = argparse.ArgumentParser(description="Build, train data for, evaluate and run the Pleiades project assistant")
    sub = parser.add_subparsers(dest="command", required=True)
    sub.add_parser("build", help="collect the project corpus and index it").set_defaults(fn=cmd_build)
    d = sub.add_parser("dataset", help="generate verified Q&A pairs and assemble the SFT set")
    d.add_argument("--only", default="", help="only chunks whose file name contains this text")
    d.add_argument("--limit", type=int, default=0)
    d.add_argument("--per-chunk", type=int, default=3)
    d.add_argument("--refusals", type=float, default=0.15)
    d.add_argument("--seed", type=int, default=42)
    d.add_argument("--assemble-only", action="store_true")
    d.set_defaults(fn=cmd_dataset)
    e = sub.add_parser("eval", help="score the assistant on the golden set")
    e.add_argument("--answers", action="store_true", help="also generate and check answers")
    e.set_defaults(fn=cmd_eval)
    c = sub.add_parser("chat", help="ask in the terminal")
    c.add_argument("question", nargs="*")
    c.set_defaults(fn=cmd_chat)
    s = sub.add_parser("serve", help="run the web app on the project knowledge base")
    s.add_argument("--port", type=int, default=8100)
    s.set_defaults(fn=cmd_serve)
    args = parser.parse_args()
    args.fn(args)


if __name__ == "__main__":
    main()
